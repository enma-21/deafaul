#!/usr/bin/env python3
"""
Recategorización masiva de publicaciones de Mercado Libre vía API.

Lee un Excel con publicaciones mal categorizadas (ID de publicación +
categoría destino) y actualiza cada una con PUT /items/{ITEM_ID}.

Requiere un access_token OAuth de Mercado Libre con permiso sobre la
cuenta vendedora (ver README.md sobre el bloqueo de DevCenter para
Venezuela y cómo obtenerlo vía una app de otro país / asesor de cuenta).

Uso típico:

    python recategorize_listings.py \
        --excel Masterbrake1937_categorizacion_v2.xlsx \
        --sheet "Confirmadas (alta confianza)" \
        --dry-run

    python recategorize_listings.py \
        --excel Masterbrake1937_categorizacion_v2.xlsx \
        --sheet "Confirmadas (alta confianza)" \
        --access-token "$ML_ACCESS_TOKEN"
"""

from __future__ import annotations

import argparse
import logging
import os
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

API_BASE = "https://api.mercadolibre.com"
DEFAULT_SHEET = "Confirmadas (alta confianza)"
RESULT_SHEET_NAME = "Resultado Actualizacion"

# Nombres de columna aceptados en el Excel, en orden de preferencia.
ITEM_ID_CANDIDATES = [
    "ID", "Item ID", "ID Publicacion", "ID Publicación",
    "MLB", "Item", "ID del item", "ID del ítem", "item_id",
]
TARGET_CATEGORY_CANDIDATES = [
    "ID Categoria Destino", "ID Categoría Destino", "Categoria Destino ID",
    "category_id_destino", "Nueva Categoria ID", "ID Categoria Nueva",
    "Category ID Destino", "category_id",
]
CURRENT_CATEGORY_LABEL_CANDIDATES = ["Categoria Actual", "Categoría actual"]
TARGET_CATEGORY_LABEL_CANDIDATES = ["Categoria Destino", "Categoría destino"]

MAX_RETRIES = 5
BASE_BACKOFF_SECONDS = 2.0
REQUEST_TIMEOUT_SECONDS = 30


@dataclass
class RowResult:
    row_number: int
    item_id: str
    target_category_id: str
    status: str  # "OK", "ERROR", "DRY_RUN", "SKIPPED"
    http_status: int | None = None
    detail: str = ""
    missing_attributes: str = ""
    timestamp: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))


def normalize(name: str) -> str:
    return " ".join(str(name).strip().lower().split())


def resolve_column(header_row: list[str], candidates: list[str], arg_value: str | None, label: str) -> int:
    """Returns the 1-based column index for a header, matching by name."""
    normalized_headers = [normalize(h) if h is not None else "" for h in header_row]

    if arg_value:
        target = normalize(arg_value)
        for idx, h in enumerate(normalized_headers, start=1):
            if h == target:
                return idx
        raise SystemExit(
            f"No se encontró la columna '{arg_value}' para {label}. "
            f"Columnas disponibles: {[h for h in header_row if h]}"
        )

    for candidate in candidates:
        target = normalize(candidate)
        for idx, h in enumerate(normalized_headers, start=1):
            if h == target:
                return idx

    raise SystemExit(
        f"No se pudo detectar automáticamente la columna de {label}. "
        f"Columnas disponibles: {[h for h in header_row if h]}. "
        f"Especifícala manualmente con el argumento correspondiente (ver --help)."
    )


def find_optional_column(header_row: list[str], candidates: list[str]) -> int | None:
    normalized_headers = [normalize(h) if h is not None else "" for h in header_row]
    for candidate in candidates:
        target = normalize(candidate)
        for idx, h in enumerate(normalized_headers, start=1):
            if h == target:
                return idx
    return None


def extract_missing_attributes(response_json: dict[str, Any]) -> str:
    """Best-effort extraction of missing/required attribute info from a ML error body."""
    causes = response_json.get("cause") or []
    if not causes:
        return ""
    parts = []
    for cause in causes:
        if not isinstance(cause, dict):
            parts.append(str(cause))
            continue
        code = cause.get("code", "")
        message = cause.get("message", "")
        references = cause.get("references") or cause.get("attribute") or ""
        parts.append(f"[{code}] {message} {references}".strip())
    return " | ".join(parts)


def update_item_category(
    session: requests.Session,
    item_id: str,
    target_category_id: str,
    access_token: str,
    extra_attributes: list[dict[str, Any]] | None = None,
) -> tuple[bool, int | None, str, str]:
    """
    Performs the PUT to update category_id (and optionally attributes).
    Returns (success, http_status, detail, missing_attributes).
    """
    url = f"{API_BASE}/items/{item_id}"
    body: dict[str, Any] = {"category_id": target_category_id}
    if extra_attributes:
        body["attributes"] = extra_attributes

    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json",
    }

    attempt = 0
    while True:
        attempt += 1
        try:
            resp = session.put(url, json=body, headers=headers, timeout=REQUEST_TIMEOUT_SECONDS)
        except requests.RequestException as exc:
            if attempt <= MAX_RETRIES:
                sleep_for = BASE_BACKOFF_SECONDS * (2 ** (attempt - 1))
                logging.warning("Error de red en %s (intento %d/%d): %s. Reintentando en %.1fs",
                                 item_id, attempt, MAX_RETRIES, exc, sleep_for)
                time.sleep(sleep_for)
                continue
            return False, None, f"Error de red tras {MAX_RETRIES} intentos: {exc}", ""

        if resp.status_code in (200, 201):
            return True, resp.status_code, "OK", ""

        if resp.status_code == 429 or resp.status_code >= 500:
            if attempt <= MAX_RETRIES:
                retry_after = resp.headers.get("Retry-After")
                sleep_for = float(retry_after) if retry_after else BASE_BACKOFF_SECONDS * (2 ** (attempt - 1))
                logging.warning("HTTP %d en %s (intento %d/%d). Reintentando en %.1fs",
                                 resp.status_code, item_id, attempt, MAX_RETRIES, sleep_for)
                time.sleep(sleep_for)
                continue

        try:
            response_json = resp.json()
        except ValueError:
            response_json = {}

        detail = response_json.get("message") or resp.text[:500]
        missing = extract_missing_attributes(response_json) if isinstance(response_json, dict) else ""
        return False, resp.status_code, detail, missing


def load_rows(excel_path: str, sheet_name: str, item_col_arg: str | None, target_col_arg: str | None):
    wb = load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(
            f"La pestaña '{sheet_name}' no existe en {excel_path}. "
            f"Pestañas disponibles: {wb.sheetnames}"
        )
    ws = wb[sheet_name]

    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    item_col = resolve_column(header_row, ITEM_ID_CANDIDATES, item_col_arg, "ID de publicación")
    target_col = resolve_column(header_row, TARGET_CATEGORY_CANDIDATES, target_col_arg, "ID de categoría destino")
    current_label_col = find_optional_column(header_row, CURRENT_CATEGORY_LABEL_CANDIDATES)
    target_label_col = find_optional_column(header_row, TARGET_CATEGORY_LABEL_CANDIDATES)

    rows = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
        item_id = row[item_col - 1].value
        target_category_id = row[target_col - 1].value
        if item_id is None or target_category_id is None:
            continue
        rows.append({
            "row_number": row_number,
            "item_id": str(item_id).strip(),
            "target_category_id": str(target_category_id).strip(),
            "current_category_label": row[current_label_col - 1].value if current_label_col else "",
            "target_category_label": row[target_label_col - 1].value if target_label_col else "",
        })

    return wb, ws, rows


def write_results(excel_path: str, results: list[RowResult]) -> str:
    """Writes a results sheet into a copy of the workbook (never overwrites the input file)."""
    wb = load_workbook(excel_path)
    if RESULT_SHEET_NAME in wb.sheetnames:
        del wb[RESULT_SHEET_NAME]
    ws = wb.create_sheet(RESULT_SHEET_NAME)

    headers = ["Fila", "Item ID", "Categoria Destino ID", "Status", "HTTP Status",
               "Detalle", "Atributos Faltantes", "Timestamp"]
    ws.append(headers)

    for r in results:
        ws.append([
            r.row_number, r.item_id, r.target_category_id, r.status,
            r.http_status, r.detail, r.missing_attributes, r.timestamp,
        ])

    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(header) + 2)

    base, ext = os.path.splitext(excel_path)
    output_path = f"{base}_resultado{ext}"
    wb.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--excel", required=True, help="Ruta al archivo .xlsx con las publicaciones a recategorizar")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"Nombre de la pestaña (default: '{DEFAULT_SHEET}')")
    parser.add_argument("--item-col", default=None, help="Nombre exacto de la columna con el ID de publicación")
    parser.add_argument("--target-col", default=None, help="Nombre exacto de la columna con el ID de categoría destino")
    parser.add_argument("--access-token", default=os.environ.get("ML_ACCESS_TOKEN"),
                         help="Access token OAuth de Mercado Libre (o variable de entorno ML_ACCESS_TOKEN)")
    parser.add_argument("--dry-run", action="store_true",
                         help="No hace ningún PUT real; solo simula y genera el reporte")
    parser.add_argument("--sleep", type=float, default=0.3,
                         help="Segundos de espera entre requests para no saturar el rate limit (default: 0.3)")
    parser.add_argument("--limit", type=int, default=None, help="Procesar solo las primeras N filas (para pruebas)")
    parser.add_argument("--log-file", default=None, help="Ruta opcional de archivo de log")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.StreamHandler(sys.stdout),
            *([logging.FileHandler(args.log_file, encoding="utf-8")] if args.log_file else []),
        ],
    )

    if not args.dry_run and not args.access_token:
        raise SystemExit(
            "Falta el access token. Pasa --access-token o define ML_ACCESS_TOKEN, "
            "o usa --dry-run para simular sin credenciales."
        )

    wb, ws, rows = load_rows(args.excel, args.sheet, args.item_col, args.target_col)
    logging.info("Publicaciones a procesar: %d (pestaña '%s')", len(rows), args.sheet)

    if args.limit:
        rows = rows[: args.limit]
        logging.info("Limitado a las primeras %d filas por --limit", args.limit)

    session = requests.Session()
    results: list[RowResult] = []
    ok_count = 0
    error_count = 0

    for i, row in enumerate(rows, start=1):
        item_id = row["item_id"]
        target_category_id = row["target_category_id"]

        if args.dry_run:
            logging.info("[DRY-RUN] (%d/%d) %s -> %s (%s)", i, len(rows), item_id,
                         target_category_id, row["target_category_label"])
            results.append(RowResult(row["row_number"], item_id, target_category_id, "DRY_RUN"))
            continue

        success, http_status, detail, missing = update_item_category(
            session, item_id, target_category_id, args.access_token
        )

        if success:
            ok_count += 1
            logging.info("(%d/%d) OK %s -> %s", i, len(rows), item_id, target_category_id)
            results.append(RowResult(row["row_number"], item_id, target_category_id, "OK", http_status))
        else:
            error_count += 1
            logging.error("(%d/%d) ERROR %s -> %s | HTTP %s | %s | Atributos faltantes: %s",
                          i, len(rows), item_id, target_category_id, http_status, detail, missing)
            results.append(RowResult(row["row_number"], item_id, target_category_id, "ERROR",
                                      http_status, detail, missing))

        time.sleep(args.sleep)

    output_path = write_results(args.excel, results)
    logging.info("Reporte guardado en: %s", output_path)
    if not args.dry_run:
        logging.info("Resumen: %d OK, %d con error, de %d procesadas", ok_count, error_count, len(rows))
        if error_count:
            logging.warning(
                "Revisa la columna 'Atributos Faltantes' del reporte para las publicaciones que "
                "requieren datos adicionales (la categoría nueva pide un atributo que la anterior no pedía)."
            )

    return 0


if __name__ == "__main__":
    sys.exit(main())
